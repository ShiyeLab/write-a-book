"""生成原创教学二进制文件；不是外部材料或真实研究。需可选依赖及 xlwt。"""
import argparse
import re
import zipfile
from datetime import datetime
from pathlib import Path

def normalize_zip(path):
    with zipfile.ZipFile(path) as z:items={n:z.read(n) for n in z.namelist()}
    with zipfile.ZipFile(path,'w',compression=zipfile.ZIP_DEFLATED) as z:
        for name,data in sorted(items.items()):
            if name=='docProps/core.xml':
                data=re.sub(rb'(<dcterms:(?:created|modified)[^>]*>)[^<]*(</dcterms:(?:created|modified)>)',rb'\g<1>2020-01-01T00:00:00Z\2',data)
            info=zipfile.ZipInfo(name,date_time=(2020,1,1,0,0,0));info.compress_type=zipfile.ZIP_DEFLATED;z.writestr(info,data)

def generate(out):
    import fitz
    from docx import Document
    from docx.oxml import OxmlElement,parse_xml
    from docx.oxml.ns import qn
    from pptx import Presentation
    from pptx.util import Inches
    from openpyxl import Workbook
    import xlwt
    out.mkdir(parents=True,exist_ok=False)
    d=Document();d.core_properties.author='Kairos';d.core_properties.last_modified_by='Kairos'
    d.add_paragraph('教学模拟：社区种子交换登记。先记录来源，再清点数量。')
    t=d.add_table(rows=2,cols=2);t.cell(0,0).text='名称';t.cell(0,1).text='数量';t.cell(1,0).text='菜豆';t.cell(1,1).text='12包'
    d.add_paragraph('登记结束后核对批次。未登记的种子暂不发放。')
    d.save(out/'registration.docx');normalize_zip(out/'registration.docx')
    complex_doc=Document();complex_doc.core_properties.author='Kairos';complex_doc.core_properties.last_modified_by='Kairos'
    p=complex_doc.add_paragraph('教学模拟：等待时间为')
    for tag,txt,texttag in [('del','20分钟','delText'),('ins','30分钟','t')]:
        e=OxmlElement('w:'+tag);e.set(qn('w:id'),'1');r=OxmlElement('w:r');t=OxmlElement('w:'+texttag);t.text=txt;r.append(t);e.append(r);p._p.append(e)
    p=complex_doc.add_paragraph('复杂对象见报告。')
    box=parse_xml('<w:pict xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:v="urn:schemas-microsoft-com:vml"><v:shape id="TeachingBox" style="width:200pt;height:40pt"><v:textbox><w:txbxContent><w:p><w:r><w:t>文本框里的教学提示</w:t></w:r></w:p></w:txbxContent></v:textbox></v:shape></w:pict>')
    p.add_run()._r.append(box)
    table=complex_doc.add_table(rows=1,cols=1);table.cell(0,0).text='外层';nested=table.cell(0,0).add_table(rows=1,cols=1);nested.cell(0,0).text='内层'
    complex_doc.save(out/'complex.docx');normalize_zip(out/'complex.docx')
    doc=fitz.open();p=doc.new_page();p.insert_text((50,70),'教学模拟：种子交换记录\n登记日期和种子批次，保留原始记录。',fontname='china-s',fontsize=20)
    scan=fitz.open();sp=scan.new_page();sp.insert_text((50,70),'教学模拟\n种子盒贴批次编号。\n不要把发芽率当作保证。',fontname='china-s',fontsize=20)
    png=sp.get_pixmap(dpi=180).tobytes('png');p=doc.new_page();p.insert_image(p.rect,stream=png);doc.new_page();doc.set_metadata({'title':'教学混合PDF','author':'Kairos'});doc.save(out/'mixed.pdf',no_new_id=True);doc.close();scan.close()
    prs=Presentation();s=prs.slides.add_slide(prs.slide_layouts[6]);s.shapes.add_textbox(Inches(1),Inches(1),Inches(6),Inches(1)).text='教学模拟：先登记批次'
    tb=s.shapes.add_table(2,2,Inches(1),Inches(2),Inches(6),Inches(1)).table;tb.cell(0,0).text='名称';tb.cell(0,1).text='数量';tb.cell(1,0).text='菜豆';tb.cell(1,1).text='12'
    prs.core_properties.author='Kairos';prs.core_properties.last_modified_by='Kairos';prs.save(out/'brief.pptx');normalize_zip(out/'brief.pptx')
    wb=Workbook();ws=wb.active;ws.title='清点';ws.append(['名称','数量']);ws.append(['菜豆',12]);ws['B3']='=B2+1';wb.properties.creator='Kairos';wb.properties.lastModifiedBy='Kairos';wb.properties.created=datetime(2020,1,1);wb.properties.modified=datetime(2020,1,1);wb.save(out/'counts.xlsx');normalize_zip(out/'counts.xlsx')
    wb=xlwt.Workbook();ws=wb.add_sheet('清点');ws.write(0,0,'教学模拟');ws.write(1,0,'菜豆');ws.write(1,1,12);wb.save(str(out/'legacy.xls'))

if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--output',required=True,type=Path);a=p.parse_args();generate(a.output)
