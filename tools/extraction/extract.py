"""可选结构化文件提取；输出文本与逐对象覆盖记录，不修改原件。"""
import argparse
import hashlib
import json
import sys
import zipfile
from datetime import date,datetime,time,timedelta
from pathlib import Path
from xml.etree import ElementTree as ET

W='{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
M='{http://schemas.openxmlformats.org/officeDocument/2006/math}'

def sha(path):
    h=hashlib.sha256()
    with Path(path).open('rb') as f:
        for block in iter(lambda:f.read(1024*1024),b''):h.update(block)
    return h.hexdigest()

def save(path,value):
    Path(path).write_text(json.dumps(value,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')

def status(items):
    if items and all(x['status']=='success' for x in items):return 'success'
    return 'partial' if any(x.get('text','').strip() or x['status'] in ['success','partial'] for x in items) else 'failed'

def code(value):return {'success':0,'partial':3,'failed':1}[value]

def item(position,text,state=None,**extra):
    return {'position':position,'status':state or ('success' if text.strip() else 'empty_needs_review'),'text':text,**extra}

def math_count(element):
    return sum(1 for e in element.iter() if e.tag==M+'oMath' or
               (e.tag==M+'oMathPara' and next(e.iter(M+'oMath'),None) is None))

def docx(path,revision):
    with zipfile.ZipFile(path) as z:
        body=ET.fromstring(z.read('word/document.xml')).find(W+'body')
        revisions=sum(1 for e in body.iter() if e.tag in {W+x for x in ['ins','del','moveFrom','moveTo']})
        if revisions and revision=='reject':raise ValueError('revisions_require_view')
        unsupported={name:len(list(body.iter(W+name))) for name in ['txbxContent','drawing','pict','footnoteReference','endnoteReference','sdt','altChunk','rPrChange','tblPrChange']}
        unsupported={k:v for k,v in unsupported.items() if v}
        math_ranges=[{'position':f'document.xml/body/{i}','kind':{W+'p':'paragraph',W+'tbl':'table'}.get(e.tag,e.tag.rsplit('}',1)[-1]),
                      'office_math':math_count(e)} for i,e in enumerate(body,1) if math_count(e)]
        if math_ranges:unsupported['office_math']=sum(r['office_math'] for r in math_ranges)
        other_parts=[n for n in z.namelist() if n.startswith(('word/header','word/footer','word/footnotes','word/endnotes')) and n.endswith('.xml')]
        if other_parts:unsupported['headers_footnotes_other_parts']=len(other_parts)
        nested=sum(len(list(t.iter(W+'tbl')))-1 for t in body.findall(W+'tbl'))
        if nested:unsupported['nested_tables']=nested
        excluded={W+'txbxContent',W+'drawing',W+'pict',M+'oMath',M+'oMathPara'}
        excluded|={W+x for x in (['del','moveFrom'] if revision=='accepted' else ['ins','moveTo'])}
        def text(e):
            if e.tag in excluded:return ''
            if e.tag in {W+'t',W+'delText'}:return e.text or ''
            if e.tag==W+'tab':return '\t'
            if e.tag in {W+'br',W+'cr'}:return '\n'
            return ''.join(text(c) for c in e)
        rows=[]
        def append_block(element,value,kind):
            row=item('body/'+str(len(rows)+1),value,kind=kind)
            count=math_count(element)
            if count:row.update(status='partial' if value.strip() else 'not_processed',unsupported={'office_math':count})
            rows.append(row)
        def blocks(parent):
            for e in parent:
                if e.tag in excluded:continue
                if e.tag==W+'p':append_block(e,text(e),'paragraph')
                elif e.tag==W+'tbl':
                    cells=[['\n'.join(text(p) for p in c.findall(W+'p')) for c in row.findall(W+'tc')] for row in e.findall(W+'tr')]
                    value='\n'.join(' | '.join(row) for row in cells) if any(c.strip() for row in cells for c in row) else ''
                    append_block(e,value,'table')
                elif e.tag in {W+'ins',W+'del',W+'moveFrom',W+'moveTo'}:blocks(e)
        blocks(body)
        # 空段落是已处理的排版间隔，不等同于整份空文件成功。
        for r in rows:
            if not r['text'] and r['status']=='empty_needs_review':r['status']='success'
        return rows,unsupported,{'revision_view':revision,'revision_objects':revisions,'unsupported_ranges':math_ranges,'known_limits':['仅正文层段落和顶层表格；不还原版面、嵌套表格、文本框、脚注、数学公式及图像内文字']}

def pdf(path):
    import fitz
    rows=[]
    with fitz.open(path) as doc:
        for i,p in enumerate(doc):
            try:
                t=p.get_text(sort=True);images=len(p.get_images())
                st='success' if t.strip() else ('needs_ocr' if images else 'empty_needs_review')
                if t.strip() and images:st='partial' # 图像区域仍须确认是否含文字
                rows.append(item('page/'+str(i+1),t,st,images=images,printed_page=p.get_label()))
            except Exception:rows.append(item('page/'+str(i+1),'','failed'))
    return rows,{}, {'known_limits':['逐页文字层，不是版面复刻；无文字不自动判定空白，混合图像区域需核对']}

def pptx(path):
    from pptx import Presentation
    rows=[];unsupported={}
    for i,slide in enumerate(Presentation(path).slides,1):
        if slide.has_notes_slide:unsupported['notes']=unsupported.get('notes',0)+1
        for j,shape in enumerate(slide.shapes,1):
            pos=f'slide/{i}/shape/{j}'
            if shape.has_table:
                rows.append(item(pos,'\n'.join(' | '.join(c.text for c in row.cells) for row in shape.table.rows),kind='table'))
            elif shape.has_text_frame:rows.append(item(pos,shape.text,kind='text'))
            else:
                rows.append(item(pos,'','not_processed',kind=str(shape.shape_type)))
                unsupported['non_text_shapes']=unsupported.get('non_text_shapes',0)+1
        if not len(slide.shapes):rows.append(item(f'slide/{i}','','empty_needs_review'))
    return rows,unsupported,{'known_limits':['按对象顺序，不推断阅读顺序；组合对象、图像文字、备注和图表未提取']}

def cached_value_fields(value):
    if isinstance(value,datetime):return {'cached_value':value.isoformat(),'cached_type':'datetime'}
    if isinstance(value,date):return {'cached_value':value.isoformat(),'cached_type':'date'}
    if isinstance(value,time):return {'cached_value':value.isoformat(),'cached_type':'time'}
    if isinstance(value,timedelta):
        return {'cached_value':{'days':value.days,'seconds':value.seconds,'microseconds':value.microseconds},'cached_type':'timedelta'}
    if value is None or isinstance(value,(str,int,float,bool)):return {'cached_value':value}
    raise TypeError('unsupported_cached_value_type')

def xlsx(path):
    from openpyxl import load_workbook
    formula=load_workbook(path,read_only=False,data_only=False);values=load_workbook(path,read_only=False,data_only=True)
    rows=[];unsupported={}
    try:
        for ws in formula:
            if ws._images or ws._charts:unsupported[ws.title]=['images_or_charts_not_processed']
            nonempty=0
            for row in ws:
                for c in row:
                    if c.value is None:continue
                    nonempty+=1;cached=values[ws.title][c.coordinate].value
                    text=str(c.value)
                    cache=cached_value_fields(cached if c.data_type=='f' else None)
                    rows.append(item(ws.title+'!'+c.coordinate,text,formula=c.value if c.data_type=='f' else None,**cache))
            if not nonempty:rows.append(item(ws.title,'','empty_needs_review'))
    finally:formula.close();values.close()
    return rows,unsupported,{'known_limits':['不计算公式；缓存缺失保留 null；不提取图表和图片文字，版式合并区域需对原件核对']}

def xls(path):
    import xlrd
    book=xlrd.open_workbook(path);rows=[]
    for ws in book.sheets():
        for i in range(ws.nrows):
            for j in range(ws.ncols):
                v=ws.cell_value(i,j)
                if v!='':rows.append(item(f'{ws.name}!R{i+1}C{j+1}',str(v),cell_type=ws.cell_type(i,j)))
        if not ws.nrows:rows.append(item(ws.name,'','empty_needs_review'))
    book.release_resources()
    return rows,{}, {'known_limits':['XLS 仅保存的单元格值；公式表达式、图片、图表和其他对象不在覆盖范围，必须原件复核']}

def extract(path,revision='reject'):
    fn={'.docx':lambda p:docx(p,revision),'.pdf':pdf,'.pptx':pptx,'.xlsx':xlsx,'.xls':xls}.get(path.suffix.lower())
    if not fn:return [],{}, {'error':'unsupported_format'}
    rows,unsupported,details=fn(path)
    return rows,unsupported,details

def save_progress(output,records):
    # 先完整写入临时文件，写入失败时保留上一份进度。
    temporary=output/'report.tmp'
    try:
        save(temporary,{'status':status(records),'records':records})
        temporary.replace(output/'report.json')
    finally:
        if temporary.exists():temporary.unlink()

def main():
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('inputs',nargs='+',type=Path);p.add_argument('--output',type=Path,required=True)
    p.add_argument('--revision-view',choices=['reject','original','accepted'],default='reject');a=p.parse_args()
    if a.output.exists() or any(not f.is_file() or a.output.resolve()==f.parent.resolve() for f in a.inputs):p.error('输入必须是文件，输出须为独立新目录')
    a.output.mkdir(parents=True);records=[{'input_name':f.name,'status':'not_processed'} for f in a.inputs];report_pending=False
    for i,source in enumerate(a.inputs,1):
        print(f'{i}/{len(a.inputs)}',flush=True)
        base={'input_name':source.name,'revision_view':a.revision_view};phase='extract'
        try:
            base['input_sha256']=sha(source);rec=dict(base)
            rows,unsupported,details=extract(source,a.revision_view);rec.update(details)
            text='\n\n'.join('<!-- '+r['position']+' ['+r['status']+'] -->\n'+r['text'] for r in rows)
            state=status(rows)
            if not any(r['text'].strip() for r in rows):state='failed'
            if unsupported and state=='success':state='partial'
            name=f'{i:04}.md';(a.output/name).write_text(text+'\n',encoding='utf-8')
            rec.update(status=state,text_file=name,objects=rows,unsupported=unsupported)
            phase='report';records[i-1]=rec;save_progress(a.output,records);report_pending=False
            continue
        except ImportError as e:rec={**base,'status':'failed','error':'missing_dependency','dependency':e.name,'install_hint':'按 docs/optional-tools.md 安装该格式的可选依赖'}
        except Exception as e:
            reason='report_write_failed' if phase=='report' else 'extraction_failed'
            if phase!='report' and isinstance(e,ValueError):reason='revisions_require_view' if str(e)=='revisions_require_view' else 'invalid_document'
            rec={**base,'status':'failed','error':reason}
        records[i-1]=rec
        try:save_progress(a.output,records);report_pending=False
        except Exception:
            report_pending=True;print(f'{i}: report_write_failed',file=sys.stderr)
    return 1 if report_pending else code(status(records))
if __name__=='__main__':sys.exit(main())
